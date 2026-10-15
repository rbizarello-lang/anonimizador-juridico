from __future__ import annotations

import logging
from dataclasses import dataclass
from pathlib import Path

from .geography import serialize_geography
from anonimizador_juridico import ExtractedPage, extract_pdf_pages, pages_to_markdown

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".txt", ".md"}
_LOG = logging.getLogger(__name__)


class UnsupportedDocumentError(ValueError):
    """Formato de arquivo fora do conjunto suportado."""


@dataclass(frozen=True, slots=True)
class ExtractedDocument:
    markdown: str
    page_count: int
    extractor: str
    ocr_used: bool = False


def extract_document(path: str | Path, *, prefer_docling: bool = True) -> ExtractedDocument:
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix == ".pdf":
        extracted = _extract_pdf(source, prefer_docling=prefer_docling)
        return ExtractedDocument(
            markdown=serialize_geography(extracted.markdown),
            page_count=extracted.page_count,
            extractor=extracted.extractor,
            ocr_used=extracted.ocr_used,
        )
    if suffix in {".txt", ".md"}:
        text = _read_text_file(source)
        return ExtractedDocument(
            markdown=pages_to_markdown([ExtractedPage(number=1, text=text)]),
            page_count=1,
            extractor="txt",
        )
    if suffix == ".docx":
        text = _read_docx(source)
        return ExtractedDocument(
            markdown=pages_to_markdown([ExtractedPage(number=1, text=text)]),
            page_count=1,
            extractor="python-docx",
        )
    raise UnsupportedDocumentError(
        f"Formato nao suportado: {suffix or '(sem extensao)'}. Use PDF, DOCX, TXT ou MD."
    )


def _extract_pdf(source: Path, *, prefer_docling: bool) -> ExtractedDocument:
    if prefer_docling:
        try:
            from .docling_extract import extract_pdf_with_docling

            markdown, page_count, ocr_used = extract_pdf_with_docling(source)
            kind = "docling-ocr" if ocr_used else "docling-nativo"
            return ExtractedDocument(
                markdown=markdown,
                page_count=page_count,
                extractor=kind,
                ocr_used=ocr_used,
            )
        except Exception as exc:
            _LOG.warning("Docling falhou em %s; usando PyMuPDF. %s", source.name, exc)
            return _extract_with_pymupdf(source, extractor="pymupdf-reserva")

    return _extract_with_pymupdf(source, extractor="pymupdf")


def _extract_with_pymupdf(source: Path, *, extractor: str) -> ExtractedDocument:
    """Reserva com o mesmo OCR seletivo: pagina digitalizada nao pode sair so com carimbo."""

    from .ocr_pages import apply_page_ocr

    pages = extract_pdf_pages(source)
    markdown = pages_to_markdown(pages)
    try:
        filled = apply_page_ocr(source, markdown)
        markdown = filled.markdown
        ocr_used = filled.pages_read > 0
    except Exception:
        _LOG.exception("OCR seletivo falhou no caminho de reserva de %s", source.name)
        ocr_used = False
    return ExtractedDocument(
        markdown=markdown,
        page_count=len(pages),
        extractor=extractor,
        ocr_used=ocr_used,
    )


def _read_text_file(path: Path) -> str:
    encodings = ("utf-8", "utf-8-sig", "latin-1", "cp1252")
    last_error: UnicodeDecodeError | None = None
    for encoding in encodings:
        try:
            return path.read_text(encoding=encoding).replace("\r\n", "\n").rstrip()
        except UnicodeDecodeError as exc:
            last_error = exc
    raise UnicodeDecodeError(
        last_error.encoding if last_error else "utf-8",
        last_error.object if last_error else b"",
        last_error.start if last_error else 0,
        last_error.end if last_error else 0,
        "Nao foi possivel ler o arquivo TXT",
    )


def _read_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - depende do ambiente
        raise RuntimeError("python-docx nao esta instalado") from exc

    from .fast_extract import IMAGE_OMITTED, docx_image_count

    document = Document(path)
    chunks: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            chunks.append(text)
        chunks.extend([IMAGE_OMITTED] * docx_image_count(paragraph._p))
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                chunks.append(" | ".join(cells))
        chunks.extend([IMAGE_OMITTED] * docx_image_count(table._tbl))
    return "\n".join(chunks).rstrip()
