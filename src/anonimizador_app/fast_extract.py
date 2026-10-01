"""Extracao rapida e fiel: PyMuPDF por pagina, ancora de pagina e marcadores do PDF.

O texto nativo sai como esta no PDF, sem reescrita nem Markdown lossy, em
milissegundos por pagina (o Docling leva segundos por pagina em CPU e, no
lote de 25/09/2026, estourou 900 s e devolveu 45 de 251 folhas).

- Ancora pela pagina fisica: ``<!-- pagina-origem: N -->`` e a folha N do PDF,
  sempre, inclusive quando ela nao tem texto.
- Geografia sem regex no corpo: os marcadores (bookmarks) do PDF, quando
  existem, viram titulo no inicio da pagina de destino. Mencao a "evento 12" no
  meio de uma frase continua no meio da frase.
- OCR so onde nao ha texto legivel (``ocr_pages``): folha escaneada, texto
  desenhado em vetor e fonte quebrada. O resto nunca passa por OCR.
- Tabela com linhas vira tabela Markdown no lugar dela; tabela sem grade
  (SISBAJUD, debitos da CDA, listagem de eventos) sai com as colunas alinhadas;
  folha com camadas de texto sobrepostas (rastreador CNIS) vira uma tabela por
  camada, na ordem do PDF. Se a conversao perder ou duplicar palavras da
  pagina, fica o texto nativo.
- Cobertura no topo (``coverage``): folha nao lida fica marcada, nao some.
"""

from __future__ import annotations

import logging
import re
import statistics
from collections import Counter, defaultdict
from collections.abc import Callable
from pathlib import Path

from .coverage import Coverage, unread_marker
from .extractors import ExtractedDocument, UnsupportedDocumentError, _read_text_file
from .report_layouts import report_plan

FAST_EXTRACTOR_VERSION = "pymupdf-rapido|14"
# Versoes cujas folhas lidas por OCR saem iguais as desta (a |14 so mudou o texto nativo: relatorios
# SISBAJUD/RENAJUD/capa, fila 3.x E12). O cache delas serve de ponto de partida: so as folhas nativas
# sao refeitas, e o OCR (~90 min no lote AÇOCARGAS) nao roda de novo.
FAST_OCR_REUSABLE_FROM = ("pymupdf-rapido|13",)
# Folha de separacao e capa do eproc: rotulo e valor na mesma linha (fila 3.x E2). A de
# documento tem "Tipo documento:"; a de evento sem documento (intimacao, prazo, ciencia)
# tem so "PAGINA DE SEPARACAO" e campos variaveis ("Prazo:", "Status:", "Data Inicial:").
_EPROC_SEPARATOR_MARKS = ("Tipo documento:", "PÁGINA DE SEPARAÇÃO")
_EPROC_LABEL_START = re.compile(
    r"^(Tipo documento:|Evento:|Data:|Usuário:|Processo:|Sequência Evento:)(?=\s|$)"
)
_EPROC_LABEL_LINE = re.compile(r"^[^:]{1,40}:$")
# Valor ao lado do rotulo (capa) ou abaixo dele, ate o proximo rotulo (separacao: +14 pt;
# "Suspensoes e Feriados" tem duas linhas).
_EPROC_SAME_LINE = 2.0
_EPROC_VALUE_GAP = 45.0
# Quadro da CDA (fila 3.x E9): rotulo em cima, valor embaixo. Rotulo -> 2a linha do rotulo.
_BOX_LABELS = {
    "Processo Administrativo": "",
    "Valor Total Inscrito Originária": "",
    "Valor Total Inscrito em": "UFIR (Lei 8383/91)",
    "Nº do agrupamento de inscrições": "",
    "Página": "",
}
_BOX_VALUE_GAP = 40.0
# Tabela so em pagina com texto e com linhas retas suficientes para uma grade.
TABLE_MIN_TEXT = 150
TABLE_MIN_SEGMENTS = 8
# A grade manda na folha se cobrir ao menos metade das palavras; menos que isso, o layout decide.
RULED_MIN_COVERAGE = 0.5
# Folha com camadas de texto sobrepostas (varias folhas impressas numa so): tabela por camada.
OVERLAP_MAX_RATIO = 0.05
# Tabela sem grade: fracao de linhas visuais com 3+ colunas; coluna = vao > 1,5 x altura da letra.
LAYOUT_MIN_LINES = 6
LAYOUT_MIN_RATIO = 0.30
LAYOUT_GAP = 1.5
_WORD = re.compile(r"\w+")
_LOG = logging.getLogger(__name__)

PageReader = Callable[[object], str]


def extract_fast(
    path: str | Path, *, read_page: PageReader | None = None, previous: str | None = None
) -> ExtractedDocument:
    """``read_page`` troca o OCR (teste); o padrao e o RapidOCR local.

    ``previous``: Markdown de uma extracao anterior do mesmo arquivo, de versao listada em
    ``FAST_OCR_REUSABLE_FROM``; a folha que ele leu por OCR e que continua pedindo OCR e copiada
    dele em vez de lida de novo."""

    source = Path(path)
    suffix = source.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(source, read_page=read_page, previous=previous)
    if suffix == ".docx":
        return ExtractedDocument(
            markdown=docx_to_markdown(source),
            page_count=1,
            extractor="python-docx",
        )
    if suffix in {".txt", ".md"}:
        return ExtractedDocument(
            markdown=_read_text_file(source) + "\n",
            page_count=1,
            extractor="txt",
        )
    raise UnsupportedDocumentError(
        f"Formato nao suportado: {suffix or '(sem extensao)'}. Use PDF, DOCX, TXT ou MD."
    )


_PAGE_ANCHOR_SPLIT = re.compile(r"(?=<!-- pagina-origem: \d+ -->\n)")
_PAGE_ANCHOR_NUMBER = re.compile(r"<!-- pagina-origem: (\d+) -->")


def _ocr_pages_of(markdown: str | None) -> dict[int, str]:
    """Folhas de uma extracao anterior lidas por OCR, inteiras (ancora e titulos inclusos)."""

    if not markdown:
        return {}
    pages: dict[int, str] = {}
    for segment in _PAGE_ANCHOR_SPLIT.split(markdown):
        found = _PAGE_ANCHOR_NUMBER.match(segment)
        if found is not None and "<!-- ocr: " in segment:
            pages[int(found.group(1))] = segment
    return pages


def _extract_pdf(
    source: Path, *, read_page: PageReader | None = None, previous: str | None = None
) -> ExtractedDocument:
    try:
        import pymupdf
    except ImportError as exc:  # pragma: no cover - depende do ambiente
        raise RuntimeError("PyMuPDF nao esta instalado") from exc

    from .ocr_pages import OCR_BROKEN_FONT, page_ocr_reason, recognize_page
    from .watermark import without_watermark

    reader = read_page or recognize_page
    parts: list[str] = []
    counts: Counter[str] = Counter()
    unread: list[int] = []
    with pymupdf.open(source) as document:
        page_count = document.page_count
        headings = _bookmarks_by_page(document.get_toc(simple=True), page_count)
        reusable = _ocr_pages_of(previous)
        for number, pdf_page in enumerate(document, start=1):
            # Selo de editor de PDF (fila 3.x E10) fora do texto, das tabelas e do OCR.
            page = without_watermark(pdf_page)
            native = page.get_text("text").replace("\r\n", "\n").strip()
            reason = page_ocr_reason(page, native)
            if reason is not None and number in reusable:
                # Folha ja lida por OCR numa versao compativel: o OCR nao roda de novo.
                parts.append(reusable[number])
                counts[reason] += 1
                continue
            parts.append(f"<!-- pagina-origem: {number} -->\n")
            for level, title in headings.get(number, ()):
                parts.append(f"{'#' * min(level + 1, 6)} {title}\n\n")
            if reason is None:
                counts["nativa"] += 1
                body = _native_body(page, native)
            else:
                body, status = _ocr_body(page, native, reason, reader, source.name, number)
                counts[status] += 1
                if status == "nao-lida":
                    unread.append(number)
            if body:
                parts.append(body + "\n\n")
    coverage = Coverage(
        pages=page_count,
        native=counts["nativa"],
        scanned=counts["escaneada"],
        vector=counts["vetorial"],
        broken_font=counts[OCR_BROKEN_FONT],
        blank=counts["em-branco"],
        unread=tuple(unread),
    )
    header = [coverage.comment()]
    if coverage.ocr:
        # Mantido para o cache e a auditoria que ja liam esta contagem.
        header.append(f"<!-- ocr-paginas: {coverage.ocr} -->")
    ocr_used = coverage.ocr > 0
    return ExtractedDocument(
        markdown="\n".join(header) + "\n" + "".join(parts),
        page_count=page_count,
        extractor="pymupdf-ocr" if ocr_used else "pymupdf",
        ocr_used=ocr_used,
    )


def _ocr_body(
    page: object,
    native: str,
    reason: str,
    reader: PageReader,
    name: str,
    number: int,
) -> tuple[str, str]:
    """Texto da pagina lida por OCR e o estado dela para a cobertura."""

    from .ocr_pages import MIN_NATIVE_CHARS, OCR_BROKEN_FONT, page_is_blank

    try:
        text = reader(page).strip()
    except Exception:  # noqa: BLE001 - falha vira pagina nao lida, nunca silencio
        _LOG.exception("OCR falhou na pagina %s de %s", number, name)
        text = None
    if text is not None and len("".join(text.split())) >= MIN_NATIVE_CHARS:
        marked = f"<!-- ocr: {reason} -->\n{text}"
        if reason == OCR_BROKEN_FONT or not native:
            # Fonte quebrada: a camada nativa e ilegivel; o OCR substitui.
            return marked, reason
        # Scan e vetor: o carimbo nativo (exato) fica, seguido da leitura.
        return f"{native}\n\n{marked}", reason
    if text is not None and reason != OCR_BROKEN_FONT and page_is_blank(page):
        return native, "em-branco"
    _LOG.warning("Pagina %s de %s sem leitura (%s)", number, name, reason)
    return "\n\n".join(item for item in (native, unread_marker(reason)) if item), "nao-lida"


def _native_body(page: object, native: str) -> str:
    """Texto nativo; tabela com grade sai em Markdown, tabela sem grade em colunas alinhadas."""

    metadata = _eproc_metadata_body(page, native)
    if metadata is not None:
        return metadata
    box = _label_box_body(page, native)
    if box is not None:
        return box
    if len("".join(native.split())) < TABLE_MIN_TEXT:
        return native
    # Extrato SISBAJUD, lista do RENAJUD, capa do processo (fila 3.x E12): rotulo com o valor.
    report = report_plan(page, native)
    if report is not None and report.tables:
        body = _text_with_tables(page, [], report)
        if _keeps_words(native, body):
            return body
        report = None
    if _overlapping_words(page) > OVERLAP_MAX_RATIO:
        # Camadas de texto sobrepostas (varias folhas impressas uma sobre a outra):
        # por posicao elas se entrelacam; na ordem do PDF cada camada sai inteira,
        # e cada registro vira uma linha (a celula quebrada volta a ficar junta).
        rows = _stream_rows(page)
        return rows if rows and _keeps_words(native, rows) else native
    ruled: str | None = None
    covered = 0.0
    boxes: list = []
    if _axis_segments(page) >= TABLE_MIN_SEGMENTS:
        planner = (lambda excluded: report_plan(page, native, exclude=excluded)) if report is not None else None
        found = _ruled_tables_body(page, planner)
        if found is not None and _keeps_words(native, found[0]):
            ruled, covered, boxes = found
    # Grade que cobre a folha: tabela Markdown. Grade pequena num canto de uma folha de
    # colunas sem grade (debitos da CDA) ou grade falsa, recusada acima (rastreador do
    # CAGED): colunas alinhadas. As colunas contam so fora das grades aceitas. Relatorio
    # reconhecido (capa do processo): as tabelas de grade ficam e os pares entram nas linhas soltas.
    if ruled is not None and (covered >= RULED_MIN_COVERAGE or report is not None):
        return ruled
    if report is not None:
        body = _text_with_tables(page, [], report)
        if _keeps_words(native, body):
            return body
    layout = _layout_body(page, exclude=boxes)
    if layout is not None and _keeps_words(native, layout):
        return layout
    return ruled if ruled is not None else native


def _eproc_metadata_body(page: object, native: str) -> str | None:
    """Folha de separacao e capa do eproc: cada rotulo com o seu valor, numa linha so.

    Os rotulos ficam numa coluna e os valores noutra (abaixo do rotulo na separacao, ao lado
    na capa); na ordem do PDF saia rotulo, rotulo, ..., valor, valor e a IA tinha de adivinhar
    os pares. Cada valor pertence ao ultimo rotulo acima dele (ou na mesma altura).
    """

    if not any(mark in native for mark in _EPROC_SEPARATOR_MARKS):
        return None
    lines: list[tuple[float, str]] = []
    labels: list[tuple[float, str]] = []
    inline: dict[float, str] = {}
    for block in page.get_text("dict")["blocks"]:
        for line in block.get("lines", []):
            text = "".join(span["text"] for span in line["spans"]).strip()
            y = line["bbox"][1]
            known = _EPROC_LABEL_START.match(text)
            if _EPROC_LABEL_LINE.match(text):
                labels.append((y, text))
            elif known is not None:
                # Rotulo e valor ja na mesma linha do PDF.
                labels.append((y, known.group(1)))
                inline[y] = text[known.end() :].strip()
            elif text:
                lines.append((y, text))
    labels.sort()
    if "Evento:" not in {text for _y, text in labels}:
        return None
    values: dict[int, list[str]] = {
        index: [inline[y]] if y in inline else [] for index, (y, _label) in enumerate(labels)
    }
    rest: list[tuple[float, str]] = []
    for y, text in sorted(lines):
        owner = None
        for index, (label_y, _label) in enumerate(labels):
            if label_y <= y + _EPROC_SAME_LINE and y - label_y <= _EPROC_VALUE_GAP:
                owner = index
        if owner is None:
            rest.append((y, text))
        else:
            values[owner].append(text)
    pairs = " — ".join(
        f"{label} {' '.join(values[index])}".rstrip() for index, (_y, label) in enumerate(labels)
    )
    rest.append((labels[0][0], pairs))
    body = "\n".join(text for _y, text in sorted(rest, key=lambda item: item[0]))
    return body if _keeps_words(native, body) else None


def _label_box_body(page: object, native: str) -> str | None:
    """Quadro de rotulo sobre valor da CDA: cada rotulo com o seu valor, numa linha por fileira.

    Na camada de texto os valores vinham antes dos rotulos e o quadro (meio da folha) caia
    depois da assinatura: "R$ 1.838,40 / Processo Administrativo / Valor Total Inscrito
    Originaria". O valor pertence ao rotulo cuja coluna ele ocupa, logo abaixo, ou ao que esta
    a esquerda na mesma altura ("Pagina 1 / 7"). So o vocabulario do quadro; palavra perdida
    ou duplicada devolve None (fica o texto nativo).
    """

    if "Valor Total Inscrito" not in native:
        return None
    lines: list[tuple[float, float, float, float, str]] = []
    for block in page.get_text("dict")["blocks"]:
        for line in block.get("lines", []):
            text = "".join(span["text"] for span in line["spans"]).strip()
            if text:
                x0, y0, x1, y1 = line["bbox"]
                lines.append((y0, y1, x0, x1, text))
    lines.sort()
    labels: list[tuple[float, float, float, float, str, list[str]]] = []
    used: set[int] = set()
    for index, (y0, y1, x0, x1, text) in enumerate(lines):
        if index in used or text not in _BOX_LABELS:
            continue
        parts, bottom = [text], y1
        tail = _BOX_LABELS[text]
        if tail:
            for other, (oy0, oy1, ox0, ox1, otext) in enumerate(lines):
                if otext == tail and -6 <= oy0 - y1 <= 8 and min(x1, ox1) > max(x0, ox0):
                    parts.append(otext)
                    bottom = oy1
                    used.add(other)
                    break
        used.add(index)
        labels.append((y0, bottom, x0, x1, " ".join(parts), parts))
    pairs: list[tuple[float, str, list[str]]] = []
    for y0, bottom, x0, x1, label, parts in labels:
        value = None
        for index, (vy0, _vy1, vx0, vx1, vtext) in enumerate(lines):
            if index in used:
                continue
            below = 0 <= vy0 - bottom <= _BOX_VALUE_GAP and min(x1, vx1) > max(x0, vx0)
            beside = abs(vy0 - y0) <= 3 and 0 <= vx0 - x1 <= 60
            if below or beside:
                value = (index, vtext)
                break
        if value is None:
            return None
        used.add(value[0])
        pairs.append((y0, x0, f"{label}: {value[1]}", [*parts, value[1]]))
    # Fileiras do quadro (mesma altura do rotulo, +-15 pt) numa linha cada, colunas em ordem.
    grouped: list[list[tuple[float, float, str, list[str]]]] = []
    for pair in sorted(pairs):
        if grouped and pair[0] - grouped[-1][0][0] <= 15:
            grouped[-1].append(pair)
        else:
            grouped.append([pair])
    rows = [
        (group[0][0], [item[2] for item in sorted(group, key=lambda item: item[1])],
         [piece for item in group for piece in item[3]])
        for group in grouped
    ]
    body = native.split("\n")
    for _y, row_pairs, removed in rows:
        positions = []
        for piece in removed:
            position = next(
                (i for i, line in enumerate(body) if line is not None and line.strip() == piece), None
            )
            if position is None:
                return None
            positions.append(position)
            body[position] = None
        body[min(positions)] = " — ".join(row_pairs)
    rebuilt = "\n".join(line for line in body if line is not None)
    return rebuilt if _keeps_words(native, rebuilt) else None


def _ruled_tables_body(page: object, planner: Callable | None = None) -> tuple[str, float, list] | None:
    """Texto com as tabelas de grade em Markdown, fracao de palavras dentro delas e as caixas.

    ``planner(caixas)`` devolve o plano de pareamento do relatorio (E12) para as linhas fora
    das grades."""

    import pymupdf

    try:
        found = page.find_tables()
    except Exception:  # noqa: BLE001 - tabela e melhoria; o texto nativo basta
        _LOG.exception("Deteccao de tabela falhou; fica o texto nativo")
        return None
    candidates = [table for table in found.tables if table.row_count >= 2 and table.col_count >= 2]
    tables = []
    for table in candidates:
        region = pymupdf.Rect(table.bbox)
        # Moldura: "tabela" que contem outra (pagina impressa da web). Fica so a de dentro.
        if any(other is not table and region.contains(pymupdf.Rect(other.bbox)) for other in candidates):
            continue
        header = getattr(table, "header", None)
        names = None
        if header is not None and getattr(header, "external", False):
            # Cabecalho fora da grade (ex.: "Placa | Marca/Modelo" acima das linhas): vira a
            # primeira linha da tabela, e as linhas dele saem do texto solto.
            names = [" ".join((name or "").split()) for name in (header.names or [])]
            region |= pymupdf.Rect(header.bbox)
        rendered = _pdf_table_markdown(table, header=names)
        if rendered:
            tables.append((tuple(region), rendered))
    if not tables:
        return None
    boxes = [pymupdf.Rect(bbox) for bbox, _rendered in tables]
    words = page.get_text("words")
    inside = sum(1 for word in words if _in_boxes(word, boxes))
    plan = planner(boxes) if planner is not None else None
    return _text_with_tables(page, tables, plan), inside / max(1, len(words)), boxes


def _stream_rows(page: object) -> str:
    """Tabela por camada, na ordem do PDF, para folha de camadas sobrepostas.

    Blocos consecutivos no fluxo e na mesma faixa vertical formam um registro;
    uma camada que vem depois no PDF nunca se entrelaca com a anterior. A linha
    sem digitos com 3+ colunas e o cabecalho da camada: as posicoes dele definem
    as colunas, e cada linha de texto do registro cai na coluna dela (a celula
    quebrada em duas linhas volta a ficar inteira).
    """

    try:
        blocks = [block for block in page.get_text("dict").get("blocks", ()) if block.get("type") == 0]
    except Exception:  # noqa: BLE001
        return ""
    groups: list[list[tuple[float, float, str]]] = []
    current: list[tuple[float, float, str]] = []
    top = bottom = 0.0
    for block in blocks:
        lines = []
        for line in block.get("lines", ()):
            text = " ".join("".join(span.get("text", "") for span in line.get("spans", ())).split())
            if text:
                lines.append((line["bbox"][0], line["bbox"][1], text))
        if not lines:
            continue
        y0, y1 = block["bbox"][1], block["bbox"][3]
        middle = (y0 + y1) / 2
        if current and not (top - 2 <= y0 <= bottom + 2 or top - 2 <= middle <= bottom + 2):
            groups.append(current)
            current = []
        if not current:
            top, bottom = y0, y1
        current.extend(lines)
        bottom = max(bottom, y1)
    if current:
        groups.append(current)

    rendered: list[str] = []
    anchors: list[float] | None = None
    for lines in groups:
        own = _column_anchors(lines)
        cells = _cells(lines, own)
        numbered = [cell for cell in cells if any(char.isdigit() for char in cell)]
        # Cabecalho: 3+ colunas sem digitos; admite uma celula "rotulo: valor" colada
        # a ele ("Total de veiculos: 2" ao lado de "Placa | Marca/Modelo").
        header = not numbered or (len(cells) >= 4 and len(numbered) == 1 and ":" in numbered[0])
        if len(cells) >= 3 and header:
            anchors = own
            rendered.append("| " + " | ".join(cells) + " |")
            rendered.append("|" + "---|" * len(cells))
        elif anchors is not None and len(own) >= 2:
            rendered.append("| " + " | ".join(_cells(lines, anchors)) + " |")
        else:
            anchors = None if len(own) < 2 else anchors
            rendered.append("  ".join(cells))
    return "\n".join(rendered)


def _column_anchors(lines: list[tuple[float, float, str]]) -> list[float]:
    """Inicio de cada coluna: linhas cujo x difere em ate 8 pt sao a mesma coluna."""

    anchors: list[float] = []
    for x0 in sorted(line[0] for line in lines):
        if not anchors or x0 - anchors[-1] > 8:
            anchors.append(x0)
    return anchors


def _cells(lines: list[tuple[float, float, str]], anchors: list[float]) -> list[str]:
    columns: list[list[tuple[float, str]]] = [[] for _ in anchors]
    for x0, y0, text in lines:
        index = 0
        for position, anchor in enumerate(anchors):
            if anchor <= x0 + 8:
                index = position
        columns[index].append((y0, text))
    return [" ".join(text for _y, text in sorted(column)) for column in columns]


def _overlapping_words(page: object) -> float:
    """Fracao de palavras que se sobrepoem a outra (mais da metade da menor caixa)."""

    try:
        words = sorted(page.get_text("words"), key=lambda word: (word[1], word[0]))
    except Exception:  # noqa: BLE001
        return 0.0
    if not words:
        return 0.0
    overlapping: set[int] = set()
    for index, word in enumerate(words):
        height = max(1.0, word[3] - word[1])
        for other_index in range(index + 1, len(words)):
            other = words[other_index]
            if other[1] - word[1] > height:
                break
            width = min(word[2], other[2]) - max(word[0], other[0])
            tall = min(word[3], other[3]) - max(word[1], other[1])
            if width <= 0 or tall <= 0:
                continue
            smaller = min((word[2] - word[0]) * height, (other[2] - other[0]) * (other[3] - other[1]))
            if smaller > 0 and width * tall > 0.5 * smaller:
                overlapping.update((index, other_index))
    return len(overlapping) / len(words)


def _in_boxes(word: tuple, boxes: list) -> bool:
    import pymupdf

    center = pymupdf.Point((word[0] + word[2]) / 2, (word[1] + word[3]) / 2)
    return any(center in box for box in boxes)


def _layout_body(page: object, *, exclude: list | None = None) -> str | None:
    """Tabela sem grade (SISBAJUD, listagem de eventos, extratos impressos da web).

    So quando ao menos 30% das linhas visuais se dividem em 3+ colunas: a peca
    corrida continua no texto nativo. A folha sai com as colunas alinhadas por
    espaco, num bloco ``text`` para o alinhamento sobreviver a qualquer visualizador.
    """

    try:
        words = page.get_text("words")
    except Exception:  # noqa: BLE001
        return None
    lines, height = _visual_lines(words)
    gap = LAYOUT_GAP * height
    split = [_segments(line, gap) for line in lines]
    # A decisao olha so as linhas fora das grades ja aceitas; o desenho usa a folha toda.
    probe = [
        segments
        for line, segments in zip(lines, split)
        if not exclude or not all(_in_boxes(word, exclude) for word in line)
    ]
    if len(probe) < LAYOUT_MIN_LINES:
        return None
    columned = sum(1 for segments in probe if len(segments) >= 3)
    if columned / len(probe) < LAYOUT_MIN_RATIO:
        return None
    widths = [(word[2] - word[0]) / max(1, len(word[4])) for line in lines for word in line]
    char_width = statistics.median(widths) if widths else 5.0
    # Sem a margem esquerda comum: menos espaco em branco, mesmo alinhamento (e menos texto para o NER).
    origin = min(line[0][0] for line in lines)
    rendered: list[str] = []
    for segments in split:
        text = ""
        for segment in segments:
            column = int((segment[0][0] - origin) / char_width)
            if len(text) < column:
                text += " " * (column - len(text))
            elif text:
                text += "  "
            text += " ".join(word[4] for word in segment)
        rendered.append(text.rstrip())
    return "```text\n" + "\n".join(rendered) + "\n```"


def _visual_lines(words: list) -> tuple[list[list], float]:
    """Agrupa palavras pela altura na folha (celulas vizinhas sao blocos diferentes no PDF)."""

    if not words:
        return [], 0.0
    heights = [word[3] - word[1] for word in words if word[3] > word[1]]
    height = statistics.median(heights) if heights else 10.0
    ordered = sorted(words, key=lambda word: ((word[1] + word[3]) / 2, word[0]))
    lines: list[list] = []
    current: list = []
    center = 0.0
    for word in ordered:
        middle = (word[1] + word[3]) / 2
        if current and abs(middle - center) > 0.45 * height:
            lines.append(sorted(current, key=lambda item: item[0]))
            current = []
        current.append(word)
        center = sum((item[1] + item[3]) / 2 for item in current) / len(current)
    if current:
        lines.append(sorted(current, key=lambda item: item[0]))
    return lines, height


def _segments(line: list, gap: float) -> list[list]:
    segments = [[line[0]]]
    for word in line[1:]:
        if word[0] - segments[-1][-1][2] > gap:
            segments.append([word])
        else:
            segments[-1].append(word)
    return segments


def _axis_segments(page: object) -> int:
    """Linhas horizontais/verticais e retangulos: sem grade nao ha tabela a procurar."""

    try:
        drawings = page.get_cdrawings()
    except Exception:  # noqa: BLE001
        return 0
    segments = 0
    for drawing in drawings:
        for item in drawing.get("items", ()):
            if item[0] == "re":
                segments += 4
            elif item[0] == "l":
                start, end = item[1], item[2]
                if abs(start[0] - end[0]) < 1 or abs(start[1] - end[1]) < 1:
                    segments += 1
    return segments


def _pdf_table_markdown(table: object, *, header: list[str] | None = None) -> str:
    rows = [header] if header and any(header) else []
    for row in table.extract():
        cells = [" ".join((cell or "").split()).replace("|", "/") for cell in row]
        if any(cells):
            rows.append(cells)
    if len(rows) < 2:
        return ""
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    lines.extend("| " + " | ".join(row) + " |" for row in rows[1:])
    return "\n".join(lines)


def _text_with_tables(page: object, tables: list[tuple[tuple, str]], plan: object | None = None) -> str:
    """Linhas fora das tabelas na ordem do PDF; cada tabela entra onde a primeira linha dela estava.

    ``plan`` (``ReportPlan``, E12): linhas trocadas pelo par rotulo/valor ou tabela do relatorio e
    linhas que ja entraram em outro texto. Os indices contam todas as linhas dos blocos de texto."""

    import pymupdf

    boxes = [pymupdf.Rect(bbox) for bbox, _rendered in tables]
    chunks: list[str] = []
    emitted: set[int] = set()
    line_index = -1
    for block in page.get_text("dict").get("blocks", ()):
        if block.get("type") != 0:
            continue
        lines: list[str] = []
        for line in block.get("lines", ()):
            line_index += 1
            text = "".join(span.get("text", "") for span in line.get("spans", ()))
            if not text.strip():
                continue
            if plan is not None:
                if line_index in plan.removed:
                    continue
                text = plan.replaced.get(line_index, text)
            center = pymupdf.Rect(line["bbox"])
            point = pymupdf.Point((center.x0 + center.x1) / 2, (center.y0 + center.y1) / 2)
            inside = next((index for index, box in enumerate(boxes) if point in box), None)
            if inside is None:
                lines.append(text)
                continue
            if lines:
                chunks.append("\n".join(lines))
                lines = []
            if inside not in emitted:
                emitted.add(inside)
                chunks.append("\n" + tables[inside][1] + "\n")
        if lines:
            chunks.append("\n".join(lines))
    for index, (_bbox, rendered) in enumerate(tables):
        if index not in emitted:
            chunks.append("\n" + rendered + "\n")
    return "\n".join(chunks).strip()


def _keeps_words(native: str, body: str) -> bool:
    """A forma nova so entra se as palavras da pagina continuarem la: nenhuma a menos, nenhuma em dobro."""

    before = Counter(_WORD.findall(native.casefold()))
    after = Counter(_WORD.findall(body.casefold()))
    total = sum(before.values())
    missing = sum((before - after).values())
    extra = sum((after - before).values())
    return missing <= max(2, total // 100) and extra <= max(3, total // 10)


def _bookmarks_by_page(toc: list, page_count: int) -> dict[int, list[tuple[int, str]]]:
    headings: dict[int, list[tuple[int, str]]] = defaultdict(list)
    for item in toc or ():
        if len(item) < 3:
            continue
        level, title, page = int(item[0]), " ".join(str(item[1]).split()), int(item[2])
        if title and 1 <= page <= page_count:
            headings[page].append((max(1, level), title))
    return headings


def docx_to_markdown(path: str | Path) -> str:
    """Paragrafos e tabelas na ordem do documento; estilo Titulo vira ``#``.

    O leitor antigo punha todas as tabelas no fim do texto, longe do
    paragrafo que as cita.
    """

    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:  # pragma: no cover - depende do ambiente
        raise RuntimeError("python-docx nao esta instalado") from exc

    document = Document(str(path))
    chunks: list[str] = []
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue
            level = _heading_level(paragraph.style.name if paragraph.style is not None else "")
            chunks.append(f"{'#' * level} {text}" if level else text)
        elif tag == "tbl":
            rendered = _table_markdown(Table(child, document))
            if rendered:
                chunks.append(rendered)
    return "\n\n".join(chunks).rstrip() + "\n"


def _heading_level(style_name: str) -> int:
    name = style_name.casefold()
    for prefix in ("heading", "titulo", "título"):
        if name.startswith(prefix):
            digits = "".join(char for char in name if char.isdigit())
            return min(int(digits), 6) if digits else 1
    return 0


def _table_markdown(table) -> str:
    rows = []
    for row in table.rows:
        cells = [" ".join(cell.text.split()).replace("|", "/") for cell in row.cells]
        if any(cells):
            rows.append(cells)
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    lines.extend("| " + " | ".join(row) + " |" for row in rows[1:])
    return "\n".join(lines)
