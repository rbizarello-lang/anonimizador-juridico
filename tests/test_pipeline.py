from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from anonimizador_app.docling_extract import clean_docling_markdown, pdf_needs_ocr
from anonimizador_app.extractors import UnsupportedDocumentError, extract_document
from anonimizador_app.pipeline import run_job


CORPUS = Path(__file__).parent / "corpus" / "caso_001_processo_sintetico.md"


class ExtractorTests(unittest.TestCase):
    def test_txt_keeps_structured_identifiers(self) -> None:
        extracted = extract_document(CORPUS)
        self.assertEqual(extracted.extractor, "txt")
        self.assertIn("529.982.247-25", extracted.markdown)
        self.assertIn("<!-- pagina-origem: 1 -->", extracted.markdown)

    def test_docx_reads_paragraphs_and_tables(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "peca.docx"
            document = Document()
            document.add_paragraph("CPF 529.982.247-25")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "OAB/SP 123.456"
            table.cell(0, 1).text = "SISBAJUD"
            document.save(path)
            extracted = extract_document(path)

        self.assertEqual(extracted.extractor, "python-docx")
        self.assertIn("529.982.247-25", extracted.markdown)
        self.assertIn("OAB/SP 123.456", extracted.markdown)
        self.assertIn("SISBAJUD", extracted.markdown)

    def test_pdf_extracts_inserted_text(self) -> None:
        import pymupdf

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "peca.pdf"
            document = pymupdf.open()
            page = document.new_page()
            page.insert_text((72, 72), "CPF 529.982.247-25")
            document.save(path)
            document.close()
            extracted = extract_document(path, prefer_docling=False)

        self.assertEqual(extracted.extractor, "pymupdf")
        self.assertIn("529.982.247-25", extracted.markdown)

    def test_unsupported_extension_is_rejected(self) -> None:
        with self.assertRaises(UnsupportedDocumentError):
            extract_document(Path("arquivo.xlsx"))

    def test_native_pdf_does_not_need_ocr(self) -> None:
        import pymupdf

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "nativo.pdf"
            document = pymupdf.open()
            page = document.new_page()
            page.insert_text((72, 72), "Evento 184 - Peticao da Uniao com texto suficiente.")
            document.save(path)
            document.close()
            self.assertFalse(pdf_needs_ocr(path))

    def test_blank_page_does_not_need_ocr(self) -> None:
        import pymupdf

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "vazio.pdf"
            document = pymupdf.open()
            document.new_page()
            document.save(path)
            document.close()
            self.assertFalse(pdf_needs_ocr(path))

    def test_stamp_plus_image_needs_ocr(self) -> None:
        import pymupdf

        from anonimizador_app.docling_extract import page_needs_ocr

        self.assertTrue(page_needs_ocr(65, 0.9))
        self.assertFalse(page_needs_ocr(1701, 0.9))
        self.assertFalse(page_needs_ocr(87, 0.0))
        self.assertFalse(page_needs_ocr(65, 0.01))

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "carimbo.pdf"
            document = pymupdf.open()
            page = document.new_page()
            page.insert_text((72, 72), "Processo Evento 1 CONTRSOCIAL9 Pagina 2")
            pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 200, 280), False)
            pix.clear_with(200)
            page.insert_image(page.rect, pixmap=pix)
            document.save(path)
            document.close()
            self.assertTrue(pdf_needs_ocr(path))

    def test_tiny_icon_does_not_need_ocr(self) -> None:
        import pymupdf

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "icone.pdf"
            document = pymupdf.open()
            page = document.new_page()
            page.insert_text((72, 72), "Processo Evento 1 CONTRSOCIAL9 Pagina 2")
            pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 8, 8), False)
            pix.clear_with(255)
            page.insert_image(pymupdf.Rect(72, 120, 80, 128), pixmap=pix)
            document.save(path)
            document.close()
            self.assertFalse(pdf_needs_ocr(path))

    def test_docling_markdown_drops_image_markers(self) -> None:
        raw = "# Titulo\n\n<!-- image -->\n\n\n\nParagrafo\n"
        cleaned = clean_docling_markdown(raw)
        self.assertNotIn("<!-- image -->", cleaned)
        self.assertIn("Paragrafo", cleaned)
        self.assertNotIn("\n\n\n", cleaned)

    def test_docling_success_uses_native_label(self) -> None:
        path = Path("peca.pdf")
        with patch(
            "anonimizador_app.docling_extract.extract_pdf_with_docling",
            return_value=("# DOCUMENTO\n\nCPF 529.982.247-25\n", 1, False),
        ):
            extracted = extract_document(path)
        self.assertEqual(extracted.extractor, "docling-nativo")
        self.assertFalse(extracted.ocr_used)
        self.assertIn("529.982.247-25", extracted.markdown)

    def test_docling_ocr_label_when_pages_lack_text(self) -> None:
        path = Path("escaneado.pdf")
        with patch(
            "anonimizador_app.docling_extract.extract_pdf_with_docling",
            return_value=("# DOCUMENTO\n\ntexto ocr\n", 3, True),
        ):
            extracted = extract_document(path)
        self.assertEqual(extracted.extractor, "docling-ocr")
        self.assertTrue(extracted.ocr_used)

    def test_docling_failure_falls_back_to_pymupdf(self) -> None:
        import pymupdf

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "reserva.pdf"
            document = pymupdf.open()
            page = document.new_page()
            page.insert_text((72, 72), "CPF 529.982.247-25")
            document.save(path)
            document.close()
            with patch(
                "anonimizador_app.docling_extract.extract_pdf_with_docling",
                side_effect=RuntimeError("docling indisponivel"),
            ):
                extracted = extract_document(path)

        self.assertEqual(extracted.extractor, "pymupdf-reserva")
        self.assertIn("529.982.247-25", extracted.markdown)


class PipelineTests(unittest.TestCase):
    def test_job_separates_markdown_from_private_map(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            export_root = root / "exportaveis"
            private_root = root / "privados"
            job = run_job(
                [CORPUS],
                export_root=export_root,
                private_root=private_root,
                job_id="lote_teste",
            )

            self.assertTrue(job.outcomes[0].ok)
            markdown_path = job.outcomes[0].markdown_path
            self.assertIsNotNone(markdown_path)
            assert markdown_path is not None
            markdown = markdown_path.read_text(encoding="utf-8")
            self.assertTrue(markdown_path.is_relative_to(export_root))
            self.assertTrue(job.map_path.is_relative_to(private_root))
            self.assertFalse(job.map_path.is_relative_to(export_root))
            self.assertEqual(job.map_path.name, "mapa.json")
            self.assertTrue((job.export_dir / "LEIA-ME.txt").is_file())
            self.assertTrue((job.export_dir / "AUDITORIA.txt").is_file())
            self.assertTrue((job.map_path.parent / "LEIA-ME.txt").is_file())
            self.assertIn("REVERTER", (job.export_dir / "LEIA-ME.txt").read_text(encoding="utf-8"))
            self.assertIn("[CPF_1]", markdown)
            self.assertIn("Helena Ficticia da Silva", markdown)
            self.assertIn("SISBAJUD", markdown)
            self.assertNotIn("529.982.247-25", markdown)
            payload = job.map_path.read_text(encoding="utf-8")
            self.assertIn("NAO ENVIAR", payload)
            self.assertFalse(job.export_allowed)

    def test_batch_reuses_tokens_across_files(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            first = root / "a.txt"
            second = root / "b.txt"
            first.write_text("CPF 529.982.247-25", encoding="utf-8")
            second.write_text("Outro trecho 52998224725", encoding="utf-8")
            job = run_job(
                [first, second],
                export_root=root / "exportaveis",
                private_root=root / "privados",
                job_id="lote_consistencia",
            )
            texts = [
                outcome.markdown_path.read_text(encoding="utf-8")
                for outcome in job.outcomes
                if outcome.markdown_path is not None
            ]
            self.assertEqual(len(texts), 2)
            self.assertTrue(all("[CPF_1]" in text for text in texts))
            self.assertEqual(job.record_count, 1)

    def test_same_root_splits_export_and_private_map(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "peca.txt"
            source.write_text("CPF 529.982.247-25 e SISBAJUD.", encoding="utf-8")
            job = run_job(
                [source],
                export_root=root,
                private_root=root,
                job_id="lote_junto",
            )
            self.assertTrue(job.outcomes[0].ok)
            markdown_path = job.outcomes[0].markdown_path
            self.assertIsNotNone(markdown_path)
            assert markdown_path is not None
            self.assertEqual(markdown_path.parent, root / "exportaveis" / "lote_junto")
            self.assertEqual(job.map_path, root / "privados" / "lote_junto" / "mapa.json")
            self.assertTrue(job.map_path.is_file())
            self.assertIn("coincidem", job.export_warning)
            self.assertFalse(job.map_path.is_relative_to(job.export_dir))


if __name__ == "__main__":
    unittest.main()
