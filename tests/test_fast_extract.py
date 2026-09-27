from __future__ import annotations

import importlib.util
import tempfile
import unittest
from pathlib import Path

from anonimizador_app.coverage import UNREAD_MARK, parse_coverage
from anonimizador_app.fast_extract import _keeps_words, extract_fast

HAS_PYMUPDF = importlib.util.find_spec("pymupdf") is not None
HAS_DOCX = importlib.util.find_spec("docx") is not None

_STAMP = "Processo 0000000-00.2026.4.04.7000/PR, Evento 1, OUT{n}, Pagina 1"


def _draw_strokes(page, count: int) -> None:
    """Texto desenhado em vetor: muitos caminhos, nenhuma camada de texto."""

    import pymupdf

    for index in range(count):
        y = 60 + (index % 40) * 16
        x = 60 + (index // 40) * 250
        page.draw_rect(pymupdf.Rect(x, y, x + 200, y + 8), color=(0, 0, 0), fill=(0, 0, 0))


def _page_block(markdown: str, number: int) -> str:
    start = markdown.index(f"<!-- pagina-origem: {number} -->")
    following = markdown.find("<!-- pagina-origem:", start + 10)
    return markdown[start:] if following < 0 else markdown[start:following]


@unittest.skipUnless(HAS_PYMUPDF, "PyMuPDF nao instalado")
class FastPdfTests(unittest.TestCase):
    def _pdf(self, directory: Path) -> Path:
        import pymupdf

        path = directory / "autos.pdf"
        with pymupdf.open() as document:
            first = document.new_page()
            first.insert_text((72, 72), "Peticao inicial da Uniao.")
            second = document.new_page()
            second.insert_text((72, 72), "Conforme o evento 12, a executada foi citada.")
            document.set_toc([[1, "Evento 1 - INIC1", 1], [1, "Evento 12 - CERT1", 2]])
            document.save(path)
        return path

    def test_page_anchors_and_bookmarks_become_geography(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            extracted = extract_fast(self._pdf(Path(tmp)))

        text = extracted.markdown
        self.assertEqual(extracted.page_count, 2)
        self.assertEqual(extracted.extractor, "pymupdf")
        self.assertIn("<!-- pagina-origem: 1 -->\n## Evento 1 - INIC1", text)
        self.assertIn("<!-- pagina-origem: 2 -->\n## Evento 12 - CERT1", text)

    def test_eproc_separator_and_cover_pair_label_with_value(self) -> None:
        import pymupdf

        labels = ["Tipo documento:", "Evento:", "Data:", "Usuário:", "Processo:", "Sequência Evento:"]
        values = ["PETIÇÃO INICIAL", "DISTRIBUIÇÃO POR SORTEIO", "19/12/2025 14:59:19",
                  "81974 - FULANO FICTICIO", "0000000-00.2025.4.04.7000", "1"]
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "sep.pdf"
            with pymupdf.open() as document:
                separator = document.new_page()
                separator.insert_text((235, 40), "PÁGINA DE SEPARAÇÃO")
                separator.insert_text((235, 95), "Documento 1")
                for index, (label, value) in enumerate(zip(labels, values)):
                    y = 160 + 31 * index
                    separator.insert_text((14, y), label)  # valor uma linha abaixo do rotulo
                    separator.insert_text((17, y + 14), value)
                cover = document.new_page()
                cover.insert_text((25, 44), "Tipo documento:")  # capa: valor ao lado
                cover.insert_text((119, 44), "CAPA PROCESSO")
                cover.insert_text((73, 58), "Evento:")
                cover.insert_text((119, 58), "abertura")
                cover.insert_text((269, 410), "PROCESSO")
                document.save(path)
            text = extract_fast(path).markdown

        expected = " — ".join(f"{label} {value}" for label, value in zip(labels, values))
        self.assertIn(expected, _page_block(text, 1))
        self.assertLess(_page_block(text, 1).index("Documento 1"), _page_block(text, 1).index(expected))
        cover_block = _page_block(text, 2)
        self.assertIn("Tipo documento: CAPA PROCESSO — Evento: abertura", cover_block)
        self.assertIn("\nPROCESSO", cover_block)

    def test_eproc_event_without_document_pairs_open_fields(self) -> None:
        import pymupdf

        rows = [("Evento:", ["EXPEDIDA INTIMACAO ELETRONICA"]), ("Data:", ["09/06/2026 13:11:55"]),
                ("Sequência Evento:", ["28"]), ("Prazo:", ["5 Dias"]),
                ("Suspensões e Feriados:", ["FERIADO A: 24/06/2026", "FERIADO B: 29/06/2026"])]
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "mov.pdf"
            with pymupdf.open() as document:
                page = document.new_page()
                page.insert_text((235, 40), "PÁGINA DE SEPARAÇÃO")
                page.insert_text((255, 95), "Evento 28")
                y = 160
                for label, values in rows:
                    page.insert_text((14, y), label)
                    for offset, value in enumerate(values, start=1):
                        page.insert_text((17, y + 14 * offset), value)
                    y += 31
                document.save(path)
            block = _page_block(extract_fast(path).markdown, 1)

        self.assertIn(
            "Evento: EXPEDIDA INTIMACAO ELETRONICA — Data: 09/06/2026 13:11:55 — Sequência Evento: 28"
            " — Prazo: 5 Dias — Suspensões e Feriados: FERIADO A: 24/06/2026 FERIADO B: 29/06/2026",
            block,
        )
        self.assertLess(block.index("Evento 28\n"), block.index("Evento: EXPEDIDA"))

    def test_event_mentioned_mid_sentence_is_not_split(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            text = extract_fast(self._pdf(Path(tmp))).markdown

        self.assertIn("Conforme o evento 12, a executada foi citada.", text)


@unittest.skipUnless(HAS_PYMUPDF, "PyMuPDF nao instalado")
class FastPdfCoverageTests(unittest.TestCase):
    """Folha nativa, escaneada, vetorial, so carimbo e vetorial ilegivel no mesmo PDF."""

    def _mixed_pdf(self, directory: Path) -> Path:
        import pymupdf

        path = directory / "misto.pdf"
        with pymupdf.open() as document:
            native = document.new_page()
            native.insert_text((72, 72), "Peticao inicial da Uniao, valor da causa R$ 1.000,00.")
            scan = document.new_page()
            scan.insert_text((20, 20), _STAMP.format(n=2), fontsize=6)
            pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 200, 280), False)
            pix.clear_with(180)
            scan.insert_image(scan.rect, pixmap=pix)
            vector = document.new_page()
            vector.insert_text((20, 20), _STAMP.format(n=3), fontsize=6)
            _draw_strokes(vector, 80)
            stamp_only = document.new_page()
            stamp_only.insert_text((20, 20), _STAMP.format(n=4), fontsize=6)
            lost = document.new_page()
            lost.insert_text((20, 20), _STAMP.format(n=5), fontsize=6)
            _draw_strokes(lost, 80)
            document.set_toc(
                [
                    [1, "Evento 1 - Doc. 1 - PETICAO INICIAL", 1],
                    [1, "Evento 1 - Doc. 2 - CONTRATO SOCIAL", 2],
                ]
            )
            document.save(path)
        return path

    def test_each_page_keeps_its_anchor_and_ocr_only_where_unreadable(self) -> None:
        calls: list[int] = []
        answers = {1: "CONTRATO SOCIAL da empresa ficticia", 2: "QUARTA ALTERACAO CONTRATUAL ficticia"}

        def reader(page) -> str:
            calls.append(page.number)
            return answers.get(page.number, "")

        with tempfile.TemporaryDirectory() as tmp:
            extracted = extract_fast(self._mixed_pdf(Path(tmp)), read_page=reader)

        text = extracted.markdown
        self.assertEqual(calls, [1, 2, 4])
        self.assertEqual(extracted.extractor, "pymupdf-ocr")
        self.assertEqual(text.count("<!-- pagina-origem:"), 5)
        self.assertIn("<!-- pagina-origem: 1 -->\n## Evento 1 - Doc. 1 - PETICAO INICIAL", text)
        scan = _page_block(text, 2)
        self.assertIn(_STAMP.format(n=2), scan)
        self.assertIn("<!-- ocr: escaneada -->\nCONTRATO SOCIAL da empresa ficticia", scan)
        self.assertIn("<!-- ocr: vetorial -->\nQUARTA ALTERACAO", _page_block(text, 3))
        self.assertNotIn("<!-- ocr", _page_block(text, 4))
        lost = _page_block(text, 5)
        self.assertIn("<!-- pagina-nao-lida: vetorial -->", lost)
        self.assertIn(UNREAD_MARK, lost)
        coverage = parse_coverage(text)
        self.assertEqual(
            (coverage.pages, coverage.native, coverage.scanned, coverage.vector, coverage.unread),
            (5, 2, 1, 1, (5,)),
        )

    def test_reader_failure_is_unread_page_not_silence(self) -> None:
        def broken(_page) -> str:
            raise ValueError("modelo indisponivel")

        with tempfile.TemporaryDirectory() as tmp:
            extracted = extract_fast(self._mixed_pdf(Path(tmp)), read_page=broken)

        self.assertEqual(parse_coverage(extracted.markdown).unread, (2, 3, 5))
        self.assertEqual(extracted.extractor, "pymupdf")

    def test_broken_font_layer_is_replaced_by_ocr(self) -> None:
        import pymupdf

        garbled = "Conta de negociacao ˭ˮ˯˰˱˲˳˴ consolidada em ˬ˭/ˬˮ/ˮˬˮ˲ no valor de ˭ˮ.˯˰˱,˲˳. " * 6
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "recibo.pdf"
            with pymupdf.open() as document:
                page = document.new_page()
                page.insert_htmlbox(pymupdf.Rect(40, 40, 560, 800), garbled)
                document.save(path)
            extracted = extract_fast(
                path,
                read_page=lambda _page: "Conta de negociacao 12345678 consolidada em 01/02/2026",
            )

        text = extracted.markdown
        self.assertIn("<!-- ocr: fonte-quebrada -->\nConta de negociacao 12345678", text)
        self.assertNotIn("ˬ˭/ˬˮ", text)
        self.assertEqual(parse_coverage(text).broken_font, 1)

    def test_ruled_table_becomes_markdown_table_in_place(self) -> None:
        import pymupdf

        above = (
            "Resposta das instituicoes financeiras ao bloqueio determinado pelo juizo, "
            "conforme a ordem protocolada no sistema e as respostas relacionadas abaixo por data."
        )
        rows = [
            ("Protocolo", "Tipo de ordem", "Resultado"),
            ("29/05/2026", "Bloqueio de valores", "R$ 1.000,00"),
            ("02/06/2026", "Reiteracao", "R$ 0,00"),
        ]
        xs, ys = [72, 222, 372, 522], [200, 225, 250, 275]
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "sisbajud.pdf"
            with pymupdf.open() as document:
                page = document.new_page()
                page.insert_textbox(pymupdf.Rect(72, 90, 522, 180), above, fontsize=11)
                for y in ys:
                    page.draw_line((xs[0], y), (xs[-1], y))
                for x in xs:
                    page.draw_line((x, ys[0]), (x, ys[-1]))
                for row, y in zip(rows, ys):
                    for cell, x in zip(row, xs):
                        page.insert_text((x + 4, y + 16), cell, fontsize=10)
                page.insert_text((72, 320), "Do que, para constar, lavrou-se a presente certidao.")
                document.save(path)
            text = extract_fast(path, read_page=lambda _page: self.fail("sem OCR aqui")).markdown

        header = "| Protocolo | Tipo de ordem | Resultado |"
        self.assertIn(header, text)
        self.assertIn("| 29/05/2026 | Bloqueio de valores | R$ 1.000,00 |", text)
        self.assertLess(text.index("relacionadas abaixo"), text.index(header))
        self.assertLess(text.index(header), text.index("Do que, para constar"))

    def test_borderless_table_keeps_columns_aligned(self) -> None:
        import pymupdf

        rows = [("Protocolo", "Ordem", "Resultado", "Saldo")]
        rows += [(f"0{day}/06/2026", "Bloqueio", f"Resposta {day}", f"R$ {day},00") for day in range(1, 9)]
        prose = (
            "A executada foi citada e nao pagou nem garantiu a execucao no prazo legal, "
            "razao pela qual a exequente requer o bloqueio de ativos financeiros. " * 4
        )
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "extrato.pdf"
            with pymupdf.open() as document:
                table_page = document.new_page()
                for index, row in enumerate(rows):
                    for cell, x in zip(row, (72, 200, 330, 450)):
                        table_page.insert_text((x, 100 + index * 18), cell, fontsize=10)
                prose_page = document.new_page()
                prose_page.insert_textbox(pymupdf.Rect(72, 72, 522, 400), prose, fontsize=11)
                document.save(path)
            text = extract_fast(path, read_page=lambda _page: self.fail("sem OCR aqui")).markdown

        table_block = _page_block(text, 1)
        self.assertIn("```text", table_block)
        header = next(line for line in table_block.splitlines() if "Protocolo" in line)
        self.assertRegex(header, r"Protocolo\s{2,}Ordem\s{2,}Resultado\s{2,}Saldo")
        self.assertNotIn("```", _page_block(text, 2))

    def test_overlapping_text_layers_stay_in_pdf_order(self) -> None:
        import pymupdf

        layers = [
            [("Empregado", "Admissao", "Fonte")] + [(f"PESSOA FICTICIA {n}{i}", f"0{i}/2023", "DIRF") for i in range(1, 8)]
            for n in ("A", "B", "C")
        ]
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "sobreposto.pdf"
            with pymupdf.open() as document:
                page = document.new_page()
                for shift, rows in enumerate(layers):
                    for index, row in enumerate(rows):
                        for cell, x in zip(row, (72, 260, 420)):
                            page.insert_text((x + shift * 0.6, 100 + index * 18), cell, fontsize=10)
                document.save(path)
            text = extract_fast(path, read_page=lambda _page: self.fail("sem OCR aqui")).markdown

        self.assertNotIn("```text", text)
        self.assertNotIn("Empregado Empregado", text)
        # Uma tabela por camada, camadas em sequencia (A antes de B), registro inteiro na linha.
        self.assertIn("| Empregado | Admissao | Fonte |", text)
        self.assertIn("| PESSOA FICTICIA A1 | 01/2023 | DIRF |", text)
        self.assertLess(text.index("PESSOA FICTICIA A7"), text.index("PESSOA FICTICIA B1"))

    def test_table_conversion_that_drops_words_is_refused(self) -> None:
        self.assertTrue(_keeps_words("Valor R$ 1,00 Uniao", "| Valor | R$ 1,00 |\nUniao"))
        self.assertFalse(_keeps_words("um dois tres quatro cinco seis sete", "um dois"))

    def test_table_conversion_that_duplicates_words_is_refused(self) -> None:
        page = "Tribunal Comarca Juiz Placa Proprietario Restricao Transferencia Valor"
        self.assertFalse(_keeps_words(page, f"| {page} |\n{page}\n{page}"))


@unittest.skipUnless(HAS_DOCX, "python-docx nao instalado")
class FastDocxTests(unittest.TestCase):
    def test_table_stays_where_it_is_cited(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "peca.docx"
            document = Document()
            document.add_heading("Dos fatos", level=1)
            document.add_paragraph("Os debitos estao na tabela abaixo.")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "CDA"
            table.cell(0, 1).text = "Valor"
            table.cell(1, 0).text = "80 6 24 000001-01"
            table.cell(1, 1).text = "R$ 1.000,00"
            document.add_paragraph("Do direito.")
            document.save(path)
            text = extract_fast(path).markdown

        self.assertIn("# Dos fatos", text)
        self.assertLess(text.index("tabela abaixo"), text.index("| CDA | Valor |"))
        self.assertLess(text.index("| CDA | Valor |"), text.index("Do direito."))
        self.assertIn("| 80 6 24 000001-01 | R$ 1.000,00 |", text)


class FastTextTests(unittest.TestCase):
    def test_markdown_is_kept_verbatim(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "nota.md"
            path.write_text("# Titulo\n\nTexto do evento 3.\n", encoding="utf-8")
            text = extract_fast(path).markdown

        self.assertEqual(text, "# Titulo\n\nTexto do evento 3.\n")


if __name__ == "__main__":
    unittest.main()
