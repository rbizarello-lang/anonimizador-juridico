"""Reverte um texto com tokens usando o mapa do lote e confere uma desanonimizacao feita fora (H9-j).

Uso (no PC; a saida tem nomes reais e vai para a pasta do mapa, que e privada):

    python tools\\reverter_lote.py <mapa.json> <arquivo com tokens (.md, .txt, .docx)>
    python tools\\reverter_lote.py <mapa.json> <peca_anonimizada.docx> --conferir <peca_desanonimizada.docx>

Sem ``--conferir``: grava ``<nome>.revertido.txt`` ao lado do mapa, com cada token trocado pela grafia preferida do
mapa (campo ``preferida``; sem ele, a grafia mais longa com acento), e imprime so contagens: tokens revertidos, tokens
que o mapa nao conhece, tokens com mais de uma grafia.

Com ``--conferir``: alinha as duas versoes linha a linha e diz, para cada ocorrencia de token, se o que a outra IA pos
no lugar e um original do mapa (comparado sem acento e caixa), se o texto fora dos tokens ficou igual e se sobrou
token. Na tela so saem contagens e rotulos de token, nunca o texto.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from pathlib import Path

TOKEN = re.compile(r"\[([A-Z][A-Z_]*?_\d+)(?:·G\d+)?\]")


def _fold(value: str) -> str:
    decomposed = unicodedata.normalize("NFKD", value)
    return " ".join("".join(c for c in decomposed if not unicodedata.combining(c)).casefold().split())


def read_text(path: Path) -> str:
    if path.suffix.lower() == ".docx":
        import docx  # python-docx, ja no ambiente do app

        document = docx.Document(str(path))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    return path.read_text(encoding="utf-8")


def load_map(path: Path) -> dict[str, dict]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    return {record["token"].strip("[]"): record for record in payload.get("records", [])}


def preferred(record: dict) -> str:
    if record.get("preferida"):
        return record["preferida"]
    originals = record.get("originals") or [""]
    return " ".join(max(originals, key=lambda v: (any(ord(c) > 127 for c in v), not v.isupper(), len(v))).split())


def revert(text: str, records: dict[str, dict]) -> tuple[str, dict[str, int]]:
    stats = {"revertidos": 0, "desconhecidos": 0}

    def swap(match: re.Match[str]) -> str:
        record = records.get(match.group(1))
        if record is None:
            stats["desconhecidos"] += 1
            return match.group(0)
        stats["revertidos"] += 1
        return preferred(record)

    return TOKEN.sub(swap, text), stats


def check(anonymized: str, reverted: str, records: dict[str, dict]) -> dict[str, int | list[str]]:
    """Confere a desanonimizacao feita fora, linha a linha (as duas versoes tem de ter as mesmas linhas)."""

    result: dict[str, int | list[str]] = {
        "ocorrencias": 0, "iguais_ao_mapa": 0, "diferentes_do_mapa": 0, "token_desconhecido": 0,
        "linhas_desalinhadas": 0, "linhas_sem_token_alteradas": 0, "tokens_que_sobraram": 0, "tokens_com_problema": [],
    }
    left, right = anonymized.split("\n"), reverted.split("\n")
    if len(left) != len(right):
        result["linhas_desalinhadas"] = abs(len(left) - len(right))
    problems: set[str] = set()
    for a, d in zip(left, right):
        result["tokens_que_sobraram"] += len(TOKEN.findall(d))
        tokens = TOKEN.findall(a)
        if not tokens:
            if a != d:
                result["linhas_sem_token_alteradas"] += 1
            continue
        parts = TOKEN.split(a)
        literal = parts[0::2]
        pattern = "^" + "".join(re.escape(piece) + ("(.+?)" if index < len(tokens) else "") for index, piece in enumerate(literal)) + "$"
        match = re.match(pattern, d, re.S)
        if match is None:
            result["linhas_desalinhadas"] += 1
            continue
        for token, value in zip(tokens, match.groups()):
            result["ocorrencias"] += 1
            record = records.get(token)
            if record is None:
                result["token_desconhecido"] += 1
                problems.add(token)
                continue
            known = {_fold(item) for item in (*record.get("originals", []), record.get("preferida", ""))} - {""}
            if _fold(value) in known:
                result["iguais_ao_mapa"] += 1
            else:
                result["diferentes_do_mapa"] += 1
                problems.add(token)
    result["tokens_com_problema"] = sorted(problems)
    return result


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("mapa", help="mapa.json do lote (pasta privada)")
    parser.add_argument("arquivo", help="texto com tokens (.md, .txt, .docx)")
    parser.add_argument("--conferir", help="a mesma peca desanonimizada por outra IA (.md, .txt, .docx)")
    parser.add_argument("--saida", help="arquivo de saida (padrao: ao lado do mapa)")
    args = parser.parse_args(argv)

    map_path = Path(args.mapa)
    records = load_map(map_path)
    text = read_text(Path(args.arquivo))
    if args.conferir:
        outcome = check(text, read_text(Path(args.conferir)), records)
        for key, value in outcome.items():
            print(f"{key}: {value}")
        ok = not (outcome["diferentes_do_mapa"] or outcome["token_desconhecido"] or outcome["linhas_desalinhadas"]
                  or outcome["linhas_sem_token_alteradas"] or outcome["tokens_que_sobraram"])
        print("RESULTADO:", "desanonimizacao conferida" if ok else "conferir os itens acima")
        return 0 if ok else 1
    reverted, stats = revert(text, records)
    destination = Path(args.saida) if args.saida else map_path.parent / (Path(args.arquivo).stem + ".revertido.txt")
    destination.write_text(reverted, encoding="utf-8")
    multi = sum(1 for record in records.values() if len({_fold(v) for v in record.get("originals", [])}) > 1)
    print(f"tokens revertidos: {stats['revertidos']}; desconhecidos: {stats['desconhecidos']}; "
          f"tokens do mapa com mais de uma grafia: {multi}")
    print(f"gravado (privado): {destination}")
    return 0 if not stats["desconhecidos"] else 1


if __name__ == "__main__":
    sys.exit(main())
